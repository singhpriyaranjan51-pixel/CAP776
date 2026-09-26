from pathlib import Path

import openpyxl

from calculation import (
    calculation_ABI,
    calculation_DCI,
    calculation_EI,
    calculation_PAI,
    calculation_aai,
    calculation_phAI,
    calculation_sri,
    calculation_tpi,
    calculation_tui,
)

REGISTRATION_NUMBER = "12609311"
EXPECTED_DAYS = 40

FEELING_SCALE = {
    "Stressed": 1,
    "Low": 2,
    "Neutral": 3,
    "Good": 4,
    "Excellent": 5,
}
SATISFACTION_SCALE = {
    "Very Unsatisfied": 1,
    "Unsatisfied": 2,
    "Neutral": 3,
    "Satisfied": 4,
    "Very Satisfied": 5,
}
ENERGY_SCALE = {"Low": 1, "Medium": 2, "High": 3}
FEELING_ALIASES = {"Happy": "Good", "Bad": "Low"}
SATISFACTION_ALIASES = {
    "Average": "Neutral",
    "Verage": "Neutral",
    "Good": "Satisfied",
    "Low": "Unsatisfied",
}
ENERGY_ALIASES = {
    "Average": "Medium",
    "Averge": "Medium",
    "Averrage": "Medium",
    "Avrage": "Medium",
}


def find_workbook():
    candidates = sorted(
        path
        for path in Path(__file__).resolve().parent.glob(f"{REGISTRATION_NUMBER}*.xlsx")
        if not path.name.startswith("~$")
    )
    if not candidates:
        raise FileNotFoundError(
            f"No Excel workbook starting with {REGISTRATION_NUMBER!r} was found "
            "beside main.py."
        )
    if len(candidates) > 1:
        names = ", ".join(path.name for path in candidates)
        raise RuntimeError(f"More than one matching workbook was found: {names}")
    return candidates[0]


def normalized_rating(value):
    if not isinstance(value, str):
        return ""
    return " ".join(value.strip().split()).title()


def main():
    workbook_path = find_workbook()
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if "Daily Log" not in workbook.sheetnames:
        raise ValueError(f"The workbook {workbook_path.name!r} has no 'Daily Log' sheet.")

    sheet = workbook["Daily Log"]
    workbook_registration = str(sheet["F2"].value or "").strip()
    if workbook_registration and workbook_registration != REGISTRATION_NUMBER:
        raise ValueError(
            f"The workbook's registration number is {workbook_registration}, "
            f"not {REGISTRATION_NUMBER}."
        )

    header_columns = {
        str(cell.value).strip(): cell.column - 1
        for cell in sheet[5]
        if cell.value is not None
    }
    required_headers = {
        "Sleep (min)",
        "Fitness (min)",
        "Study (min)",
        "Coding (min)",
        "Class (min)",
        "Total Tracked (min)",
        "Free/Unaccounted (min)",
        "Day's Feeling",
        "Satisfaction Level",
        "Energy Level",
    }
    missing_headers = required_headers - header_columns.keys()
    if missing_headers:
        raise ValueError(f"The Daily Log is missing columns: {', '.join(sorted(missing_headers))}")

    values = {name: [] for name in required_headers}
    valid_days = 0
    invalid_rating_rows = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        if not row or row[0] is None:
            continue
        valid_days += 1
        for header in required_headers - {"Day's Feeling", "Satisfaction Level", "Energy Level"}:
            value = row[header_columns[header]]
            if isinstance(value, (int, float)):
                values[header].append(value)

        feeling = normalized_rating(row[header_columns["Day's Feeling"]])
        satisfaction = normalized_rating(row[header_columns["Satisfaction Level"]])
        energy = normalized_rating(row[header_columns["Energy Level"]])
        feeling = FEELING_ALIASES.get(feeling, feeling)
        satisfaction = SATISFACTION_ALIASES.get(satisfaction, satisfaction)
        energy = ENERGY_ALIASES.get(energy, energy)
        if (
            feeling in FEELING_SCALE
            and satisfaction in SATISFACTION_SCALE
            and energy in ENERGY_SCALE
        ):
            values["Day's Feeling"].append(FEELING_SCALE[feeling])
            values["Satisfaction Level"].append(SATISFACTION_SCALE[satisfaction])
            values["Energy Level"].append(ENERGY_SCALE[energy])
        else:
            invalid_rating_rows.append(row_number)

    coding = calculation_tpi(values["Coding (min)"])
    academic = calculation_aai(values["Study (min)"], values["Class (min)"])
    fitness = calculation_phAI(values["Fitness (min)"])
    sleep = calculation_sri(values["Sleep (min)"])
    unaccounted = calculation_ABI(values["Free/Unaccounted (min)"])
    tracked = calculation_tui(values["Total Tracked (min)"])
    emotional = calculation_EI(
        values["Day's Feeling"],
        values["Satisfaction Level"],
        values["Energy Level"],
    )
    consistency = calculation_DCI(valid_days, EXPECTED_DAYS)
    overall = calculation_PAI(coding, academic, fitness, sleep, tracked, emotional, consistency)

    print(f"Registration number: {REGISTRATION_NUMBER}")
    print(f"Workbook: {workbook_path.name}")
    print(f"Daily entries: {valid_days}")
    print(f"TPI: {coding:.2f} min/day")
    print(f"AAI: {academic:.2f} min/day")
    print(f"PhAI: {fitness:.2f} min/day")
    print(f"SRI: {sleep:.2f} min/day")
    print(f"ABI: {unaccounted:.2f} min/day")
    print(f"TUI: {tracked:.2f} min/day")
    print(f"EI: {emotional:.2f}")
    print(f"DCI: {consistency:.2f}% (based on {EXPECTED_DAYS} expected days)")
    print(f"PAI: {overall:.2f}")
    if invalid_rating_rows:
        print(
            "Rows with missing or invalid feeling/satisfaction/energy ratings "
            f"(excluded from EI): {', '.join(map(str, invalid_rating_rows))}"
        )
    workbook.close()


if __name__ == "__main__":
    main()
